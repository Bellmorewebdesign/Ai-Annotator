"""Tests for regions (boxes/circles), OCR, conversion and input validation."""

from __future__ import annotations

import shutil

import pikepdf
import pytest
from conftest import load
from engine_helpers import CATS, changed_mask, fraction, user_box_to_pixels
from fixtures import ACADEMIC_TITLE, academic_docx, figure_table_pdf, scanned_pdf, simple_pdf
from PIL import Image, ImageDraw, ImageFont

from annotator.engine.annotate import RegionSelection, TextSelection, place_annotations, render_and_verify
from annotator.engine.convert import docx_to_pdf, find_soffice, image_to_pdf, text_to_pdf, txt_file_to_pdf
from annotator.engine.ocr import find_tesseract
from annotator.engine.validate import UnsupportedFile, sanitize_filename, sniff

TESSERACT = find_tesseract()
needs_ocr = pytest.mark.skipif(TESSERACT is None, reason="Tesseract not installed")


def test_figure_and_table_regions_are_detected_with_captions(tmp):
    src = tmp / "fig.pdf"
    figure_table_pdf(src)
    res, dt = load(src)
    regions = res.doc.pages[0].regions
    img = next(r for r in regions if r.kind == "image")
    assert img.confidence == "high" and img.caption.startswith("Figure 1")
    x0, y0, x1, y1 = img.user_box
    assert abs(x0 - 72) < 1 and abs(x1 - 372) < 1 and abs(y0 - 470) < 1 and abs(y1 - 650) < 1
    table = next(r for r in regions if r.kind == "table")
    assert table.confidence == "medium" and table.caption.startswith("Table 1")


def test_box_and_circle_annotations_render_on_regions(tmp):
    src, out = tmp / "fig.pdf", tmp / "out.pdf"
    figure_table_pdf(src)
    res, dt = load(src)
    regions = {r.kind: r for r in res.doc.pages[0].regions}
    placement = place_annotations(res.doc, dt, [], [
        RegionSelection("r1", regions["image"].id, "box", "main", "Key chart"),
        RegionSelection("r2", regions["table"].id, "circle", "evidence", "Main numbers"),
    ], CATS)
    outcome = render_and_verify(src, out, placement)
    assert outcome.verified and len(placement.placed) == 2
    with pikepdf.open(out) as pdf:
        subtypes = sorted(str(a.Subtype) for a in pdf.pages[0].Annots)
    assert subtypes == ["/Circle", "/Square"]
    mask = changed_mask(src, out, 0)
    crop = res.doc.pages[0].crop_box
    x0, y0, x1, y1 = regions["image"].user_box
    # The box outline lies just outside the image edge; the image interior is untouched.
    edge = user_box_to_pixels(crop, 0, (x0 - 5, y0 - 5, x0 - 1, y1 + 5))
    assert fraction(mask, edge, shrink=0) > 0.1
    inner = user_box_to_pixels(crop, 0, (x0 + 10, y0 + 10, x1 - 10, y1 - 10))
    assert fraction(mask, inner) < 0.01


def test_unreliable_region_falls_back_to_caption(tmp):
    src, out = tmp / "fig.pdf", tmp / "out.pdf"
    figure_table_pdf(src)
    res, dt = load(src)
    page = res.doc.pages[0]
    img = next(r for r in page.regions if r.kind == "image")
    img.confidence = "low"  # simulate a figure whose area could not be determined
    placement = place_annotations(res.doc, dt, [], [RegionSelection("r1", img.id, "box", "main")], CATS)
    outcome = render_and_verify(src, out, placement)
    assert outcome.verified
    assert placement.placed[0].kind == "highlight" and placement.placed[0].match == "caption"
    assert placement.placed[0].text.startswith("Figure 1")


@needs_ocr
def test_scanned_pdf_is_ocrd_and_highlights_land_on_words(tmp):
    text_src, scan, out = tmp / "text.pdf", tmp / "scan.pdf", tmp / "out.pdf"
    truth = simple_pdf(text_src)
    scanned_pdf(scan, text_src)
    res, dt = load(scan, tesseract=TESSERACT)
    page = res.doc.pages[0]
    assert page.is_ocr and page.ocr_mean_conf is not None and page.ocr_mean_conf > 80
    phrase = "choosing the right evidence"
    bid = next(b for b, (s, e) in dt.block_spans.items() if phrase in dt.text[s:e])
    placement = place_annotations(res.doc, dt, [TextSelection("s", "highlight", bid, phrase, "main")], [], CATS)
    outcome = render_and_verify(scan, out, placement, doc=res.doc)
    assert placement.placed and placement.placed[0].match == "exact"
    assert outcome.content_preserved
    # Highlight lands where the words were drawn in the original text PDF.
    words = truth.find(phrase)
    box = (words[0].x0, words[0].baseline - 2, words[-1].x1, words[0].baseline + 7)
    mask = changed_mask(scan, out, 0)
    assert fraction(mask, user_box_to_pixels(page.crop_box, 0, box)) > 0.4
    before = truth.words[truth.words.index(words[0]) - 1]
    assert fraction(mask, user_box_to_pixels(page.crop_box, 0, (before.x0 + 1, before.baseline, before.x1 - 2,
                                                                 before.baseline + 6)), shrink=0) < 0.05


def test_scanned_pdf_without_ocr_is_flagged(tmp):
    text_src, scan = tmp / "text.pdf", tmp / "scan.pdf"
    simple_pdf(text_src)
    scanned_pdf(scan, text_src)
    res, dt = load(scan, tesseract=None)
    assert res.skipped_pages == [0]
    assert any("OCR" in w for w in res.warnings)


@needs_ocr
def test_image_input_is_converted_and_ocrd(tmp):
    png, pdf = tmp / "photo.png", tmp / "photo.pdf"
    img = Image.new("RGB", (1700, 600), "white")
    d = ImageDraw.Draw(img)
    font = ImageFont.truetype(str(shutil.os.path.join(
        str(__import__("annotator").__path__[0]), "assets", "fonts", "DejaVuSerif.ttf")), 44)
    d.text((80, 120), "Photosynthesis converts light energy", font=font, fill="black")
    d.text((80, 200), "into chemical energy stored in sugar.", font=font, fill="black")
    img.save(png, dpi=(200, 200))
    assert sniff(png).kind == "png"
    image_to_pdf(png, pdf)
    res, dt = load(pdf, tesseract=TESSERACT)
    assert res.doc.pages[0].is_ocr
    assert "chemical energy" in dt.text


def test_text_file_conversion_handles_encodings(tmp):
    src = tmp / "notes.txt"
    src.write_bytes("Café résumé notes.\n\nSecond paragraph with “quotes”.".encode("cp1252"))
    s = sniff(src)
    assert s.kind == "txt" and s.text_encoding == "cp1252"
    out = tmp / "notes.pdf"
    txt_file_to_pdf(src, out, s.text_encoding)
    res, dt = load(out)
    assert "Café résumé notes." in dt.text
    assert "“quotes”" in dt.text


def _make_docx(path):
    import docx

    d = docx.Document()
    d.add_heading("Lab Report", level=1)
    d.add_paragraph("The reaction rate doubled when the temperature increased by ten degrees.")
    d.save(str(path))


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_docx_conversion_with_libreoffice(tmp):
    src, out = tmp / "r.docx", tmp / "r.pdf"
    _make_docx(src)
    assert sniff(src).kind == "docx"
    result = docx_to_pdf(src, out, find_soffice())
    assert result.method == "libreoffice"
    res, dt = load(out)
    assert "reaction rate doubled" in dt.text


def test_docx_fallback_without_libreoffice_is_flagged(tmp):
    src, out = tmp / "r.docx", tmp / "r.pdf"
    _make_docx(src)
    result = docx_to_pdf(src, out, None)
    assert result.method == "docx-text" and "simplified" in result.notes[0]
    res, dt = load(out)
    assert "reaction rate doubled" in dt.text


@pytest.mark.parametrize("content,message", [
    (b"", "empty"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, ".doc"),
    (b"\x00\x01\x02\x03" * 100, "not supported"),
    (b"GIF89a" + b"\x00" * 100, "not supported"),
])
def test_unsupported_files_are_rejected(tmp, content, message):
    f = tmp / "file.bin"
    f.write_bytes(content)
    with pytest.raises(UnsupportedFile) as info:
        sniff(f)
    assert message in str(info.value)


def test_pdf_detected_by_content_not_extension(tmp):
    f = tmp / "paper.txt"
    simple_pdf(f)
    assert sniff(f).kind == "pdf"


def test_sanitize_filename_blocks_path_tricks():
    assert sanitize_filename("../../etc/passwd") == "passwd"
    assert sanitize_filename("..\\..\\win.ini") == "win.ini"
    assert sanitize_filename("my <paper>.pdf") == "my _paper_.pdf"
    assert sanitize_filename("") == "document"
    assert len(sanitize_filename("a" * 300 + ".pdf")) <= 120


def test_text_to_pdf_rejects_empty_text(tmp):
    with pytest.raises(UnsupportedFile):
        text_to_pdf("   \n\n  ", tmp / "x.pdf")


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_libreoffice_two_column_justified_document(tmp):
    """A real-world PDF producer: justified text, ligatures, hyphenation, two columns."""
    from engine_helpers import annotate

    src_docx, src = tmp / "paper.docx", tmp / "paper.pdf"
    academic_docx(src_docx)
    docx_to_pdf(src_docx, src, find_soffice())
    res, dt = load(src)
    blocks = [b for p in res.doc.pages for b in p.blocks]
    headings = [b.text for b in blocks if b.role == "heading"]
    # The wrapped title (full width, then a short second line) must stay one heading.
    assert headings[:3] == [ACADEMIC_TITLE, "Introduction", "Methods"]
    order = [b.text.split()[0] for b in blocks]
    assert order.index("Methods") < order.index("Results") < order.index("Discussion")

    def bid(phrase):
        return next(b for b, (s, e) in dt.block_spans.items() if phrase in dt.text[s:e])

    items = [
        ("Students recalled 64 percent of quizzed material compared with 47 percent of reread material",
         bid("Students recalled"), "highlight"),
        ("the study has limitations", bid("limitations"), "underline"),
        ("our design equated exposure time", bid("Counterarguments"), "highlight"),
    ]
    placement, outcome = annotate(res.doc, dt, src, tmp / "out.pdf", items)
    assert outcome.verified and not placement.rejected, placement.rejected
    assert len(placement.specs[0].quads) >= 2  # wraps across lines in a narrow column
    # Justified word gaps must not split a line: at most one continuous quad per line.
    # (How many lines a phrase spans depends on the fonts LibreOffice substitutes.)
    for spec in placement.specs:
        tops = [round(max(q[1::2]), 1) for q in spec.quads]
        assert len(tops) == len(set(tops)), "more than one quad on the same line"
